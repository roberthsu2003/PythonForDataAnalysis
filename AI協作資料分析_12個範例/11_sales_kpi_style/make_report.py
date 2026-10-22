"""
11｜業務週報：整季達成率與格式化報表自動化腳本
"""
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
DATA_PATH = BASE_DIR / "data" / "sales_kpi.csv"
OUTPUT_FILE = BASE_DIR / "output" / "第一季業務達成率.xlsx"


def load_data(csv_path: Path) -> pd.DataFrame:
    if not csv_path.exists():
        raise FileNotFoundError(f"找不到資料檔案：{csv_path}")
    return pd.read_csv(csv_path)


def quarter_rate(df: pd.DataFrame, keys: list) -> pd.DataFrame:
    """整季達成率＝實績合計 ÷ 目標合計（不平均每週比率，避免目標不同時權重失真）"""
    result = df.groupby(keys)[["週目標", "週實績"]].sum().rename(columns={"週目標": "整季目標", "週實績": "整季實績"})
    result["整季達成率"] = result["整季實績"] / result["整季目標"]
    return result.sort_values("整季達成率", ascending=False).reset_index()


def export_excel(group_df: pd.DataFrame, person_df: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path) as writer:
        group_df.to_excel(writer, sheet_name="組別達成率", index=False)
        person_df.to_excel(writer, sheet_name="個人達成率", index=False)


def format_workbook(output_path: Path) -> None:
    """只改顯示格式，儲存格保留數值"""
    wb = load_workbook(output_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        headers = {cell.value: cell.column_letter for cell in ws[1]}
        for cell in ws[1]:
            cell.fill = PatternFill("solid", start_color="1F497D")
            cell.font = Font(bold=True, color="FFFFFF")
        for col in ["整季目標", "整季實績"]:
            for cell in ws[headers[col]][1:]:
                cell.number_format = "#,##0"
        for cell in ws[headers["整季達成率"]][1:]:
            cell.number_format = "0.0%"
            if cell.value < 1:
                cell.fill = PatternFill("solid", start_color="FCE4D6")
        for letter in headers.values():
            ws.column_dimensions[letter].width = 14
    wb.save(output_path)


def main():
    print("=" * 50)
    print("🚀 開始執行業務週報達成率腳本...")
    print("=" * 50)

    df = load_data(DATA_PATH)
    group_df = quarter_rate(df, ["組別"])
    person_df = quarter_rate(df, ["組別", "業務"])
    export_excel(group_df, person_df, OUTPUT_FILE)
    format_workbook(OUTPUT_FILE)

    print("\n📊 【各組整季達成率】（整季實績 ÷ 整季目標）")
    for _, row in group_df.iterrows():
        print(f"  - {row['組別']}：{row['整季實績']:>12,.0f} ÷ {row['整季目標']:>12,.0f} = {row['整季達成率']:.1%}")
    below = person_df[person_df["整季達成率"] < 1]
    print(f"\n⚠️ 【未達 100% 的業務】共 {len(below)} 位")
    for _, row in below.iterrows():
        print(f"  - {row['業務']}（{row['組別']}）：{row['整季達成率']:.1%}")
    print(f"\n✅ 成果已成功匯出至：{OUTPUT_FILE.relative_to(BASE_DIR)}")
    print("=" * 50)


if __name__ == "__main__":
    main()
